"""Pick the primary manuscript file, run the right parser, merge bibliographies."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from app.models.job import UploadedFile
from app.models.manuscript import ManuscriptStructure
from app.parsers import bibtex_parser, docx_parser, latex_parser, markdown_parser

PRIORITY = {".tex": 0, ".md": 1, ".markdown": 1, ".docx": 2}


@dataclass
class ParseOutcome:
    manuscript: ManuscriptStructure
    primary: Path
    bib_changes: list[str] = field(default_factory=list)
    bib_warnings: list[str] = field(default_factory=list)


def choose_primary(upload_dir: Path, files: list[UploadedFile]) -> Path:
    manuscripts = [f for f in files if f.kind == "manuscript"]
    if not manuscripts:
        raise ValueError("No manuscript file was uploaded.")
    best_rank = min(PRIORITY[f.extension] for f in manuscripts)
    candidates = [upload_dir / f.stored_name for f in manuscripts if PRIORITY[f.extension] == best_rank]
    if candidates[0].suffix == ".tex" and len(candidates) > 1:
        return latex_parser.choose_main_tex(candidates)
    return max(candidates, key=lambda p: p.stat().st_size)


def parse_primary(primary: Path, upload_dir: Path) -> ManuscriptStructure:
    ext = primary.suffix.lower()
    if ext == ".tex":
        return latex_parser.parse_latex(primary)
    if ext in {".md", ".markdown"}:
        return markdown_parser.parse_markdown(primary)
    if ext == ".docx":
        return docx_parser.parse_docx(primary, figure_dir=upload_dir)
    raise ValueError(f"Unsupported manuscript type: {ext}")


def plain_text_of(primary: Path, max_chars: int = 60000) -> str:
    ext = primary.suffix.lower()
    if ext == ".docx":
        from docx import Document

        doc = Document(str(primary))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)[:max_chars]
    return primary.read_text(encoding="utf-8", errors="replace")[:max_chars]


def merge_bibliographies(manuscript: ManuscriptStructure, primary: Path, upload_dir: Path,
                         files: list[UploadedFile]) -> tuple[list[str], list[str]]:
    """Load uploaded .bib files (preferring those the manuscript names) into references."""
    bib_files = [upload_dir / f.stored_name for f in files if f.kind == "bibliography"]
    if not bib_files:
        return [], []
    wanted: list[str] = []
    if primary.suffix == ".tex":
        wanted = latex_parser.referenced_bib_files(primary)
    elif primary.suffix in {".md", ".markdown"}:
        wanted = markdown_parser.bibliography_from_front_matter(primary)
    ordered = sorted(bib_files, key=lambda p: (p.name not in wanted, p.name))
    changes: list[str] = []
    warnings: list[str] = []
    seen: set[str] = set()
    bib_refs = []
    for path in ordered:
        result = bibtex_parser.parse_bibtex_file(path)
        changes += result.changes
        warnings += result.warnings
        for ref in result.references:
            if ref.key in seen:
                continue
            seen.add(ref.key)
            bib_refs.append(ref)
    if bib_refs:
        # A real .bib file supersedes \bibitem entries and plain-text reference lists.
        if any(r.origin == "thebibliography" for r in manuscript.references):
            warnings.append("Both \\bibitem entries and a .bib file were found; the .bib file is used.")
        manuscript.references = bib_refs
    return changes, warnings


def parse_uploads(upload_dir: Path, files: list[UploadedFile]) -> ParseOutcome:
    primary = choose_primary(upload_dir, files)
    manuscript = parse_primary(primary, upload_dir)
    changes, warnings = merge_bibliographies(manuscript, primary, upload_dir, files)
    return ParseOutcome(manuscript=manuscript, primary=primary, bib_changes=changes, bib_warnings=warnings)
