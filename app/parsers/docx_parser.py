"""Deterministic DOCX manuscript parser (python-docx).

Uses Word paragraph styles first (Title, Heading 1-4, Caption, List) and falls
back to text heuristics (numbered headings, "Abstract", "Keywords", "Figure N"
captions). Body text is converted straight to LaTeX with bold/italic runs kept.
Embedded images are extracted into the job's upload folder so the formatter can
place them. Word equations (OMML) are converted to LaTeX math: on their own line
they become numbered equations, inside a sentence inline math.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from app.models.manuscript import Figure, ManuscriptStructure, SourceInfo, Table, TextBlock
from app.parsers.omml import omml_to_latex
from app.parsers.outline import TO_SUPERSCRIPT, Outline, authors_from_lines, nested_list_latex, special_part
from app.parsers.text_utils import latex_escape, normalize_space, smart_title, strip_heading_number

NUMBERED_HEADING = re.compile(r"^((\d+(\.\d+){0,3})|([IVX]+))[.)]?\s+([A-Z][^.!?]{1,80})$")
CAPTION_RE = re.compile(r"^(fig\.?|figure|table)\s*(\d+|[ivx]+)?\s*[:.\-\u2013\u2014]\s*(.*)$", re.I)
KEYWORDS_RE = re.compile(r"^(keywords|index terms|key words)\s*[:\u2014\-]+\s*(.*)$", re.I)
INLINE_ABSTRACT_RE = re.compile(r"^abstract\s*[:.\u2014\-]+\s*(.+)$", re.I)
IMAGE_EXTENSIONS = {"image/png": ".png", "image/jpeg": ".jpg", "image/jpg": ".jpg"}
EQUATION_NUMBER_RE = re.compile(r"^#?\(?\d+(\.\d+)?\)?$")  # "(1)" typed after a display equation
RUN_CONTAINERS = {qn("w:hyperlink"), qn("w:ins"), qn("w:smartTag")}


def parse_docx(path: Path, figure_dir: Path | None = None) -> ManuscriptStructure:
    reader = _DocxReader(Document(str(path)), figure_dir or path.parent)
    reader.read()
    outline = reader.outline
    sections = outline.all_sections()
    authors = authors_from_lines(reader.author_lines)

    warnings = reader.warnings
    if reader.unsupported_images:
        warnings.append(f"{reader.unsupported_images} embedded image(s) use a format Word stores as EMF/WMF/other "
                        "and could not be extracted; export them as PNG and upload separately.")
    if not sections:
        warnings.append("No headings detected; use Word 'Heading 1/2' styles for best results.")
    if reader.footnotes:
        warnings.append(f"{len(reader.footnotes)} footnote(s) outside the body text (on the title, authors, a heading, "
                        "caption, table or reference) were left out; add them by hand.")

    confidence = 0.3
    confidence += 0.15 if reader.title else 0
    confidence += 0.15 if outline.abstract else 0
    confidence += 0.25 if len(sections) >= 2 else 0
    confidence += 0.1 if authors else 0

    return ManuscriptStructure(
        source=SourceInfo(primary_file=path.name, format="docx", confidence=min(confidence, 1.0), warnings=warnings),
        title=reader.title,
        authors=authors,
        abstract="\n\n".join(outline.abstract),
        keywords=reader.keywords,
        sections=sections,
        figures=reader.figures,
        tables=reader.tables,
        acknowledgments="\n\n".join(outline.acknowledgments) or None,
        references=outline.references,
        citation_keys_used=[],
        text_format="latex",
    )


class _DocxReader:
    """Reads a Word document paragraph by paragraph. Lines before the abstract or first heading are the
    title and author block ("front" part)."""

    def __init__(self, document, figure_dir: Path) -> None:
        self.document = document
        self.figure_dir = figure_dir
        self.outline = Outline(part="front")
        self.title = ""
        self.author_lines: list[str] = []
        self.keywords: list[str] = []
        self.figures: list[Figure] = []
        self.tables: list[Table] = []
        self.warnings: list[str] = []
        self.table_caption = ""  # a "Table N:" caption written above its table
        self.unsupported_images = 0
        self.list_block: TextBlock | None = None  # the list being read
        self.list_entries: list[tuple[int, bool, str]] = []  # its items as (level, ordered, LaTeX)
        self.footnotes = _footnotes(document)  # id -> LaTeX text, removed once placed

    def read(self) -> None:
        for child in self.document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                self.read_paragraph(Paragraph(child, self.document))
            elif child.tag == qn("w:tbl"):
                self.read_table(DocxTable(child, self.document))

    def read_table(self, table: DocxTable) -> None:
        rows = [[latex_escape(normalize_space(cell.text)) for cell in row.cells] for row in table.rows]
        if not rows:
            return
        tab = Table(id=f"tab{len(self.tables) + 1}", caption=self.table_caption, header=rows[0], rows=rows[1:])
        self.table_caption = ""
        self.tables.append(tab)
        self.outline.add(TextBlock(kind="table_ref", content=tab.id))

    def read_paragraph(self, p: Paragraph) -> None:
        text = normalize_space(p.text)
        style = _style_name(p)

        if self.read_images(p) and not text:
            return
        if self.read_display_math(p, text):
            return
        if not text or style == "subtitle":
            return
        if style == "title" and not self.title:
            self.title = latex_escape(text)
            return

        caption = CAPTION_RE.match(text)
        if caption and (style == "caption" or len(text) < 300) and self.add_caption(caption):
            return
        if keywords := KEYWORDS_RE.match(text):
            self.keywords = [k.strip().rstrip(".") for k in re.split(r"[,;]", keywords.group(2)) if k.strip()]
            return
        abstract = INLINE_ABSTRACT_RE.match(text)
        if abstract and self.outline.part in {"front", "abstract"}:
            self.outline.abstract.append(latex_escape(abstract.group(1)))
            self.outline.part = "abstract"
            return
        # A plain "Abstract" line starts the abstract anywhere; a plain "References" or "Acknowledgment"
        # line only counts after the front matter (an author could be called "References").
        part = special_part(text)
        if part and (part == "abstract" or self.outline.part != "front"):
            self.outline.part = part
            return
        if level := _heading_level(p, text):
            self.outline.heading(latex_escape(smart_title(strip_heading_number(text))), level)
            return
        if not self.title and self.outline.part == "front":
            self.title = latex_escape(text)
            return

        if self.outline.part == "front":
            self.author_lines.append(_marked_text(p) or text)
            return
        if self.outline.part == "references":
            self.outline.add_reference(latex_escape(text))
            return
        latex = _runs_to_latex(p, self.footnotes) or latex_escape(text)
        if self.outline.part == "body" and "list" in style:
            self.add_list_item(latex, style)
        else:
            self.outline.add(TextBlock(kind="paragraph", content=latex, format="latex"))

    def read_display_math(self, p: Paragraph, text: str) -> bool:
        """A paragraph holding only an equation (and perhaps its "(1)") becomes an equation block."""
        equations = p._p.findall(".//" + qn("m:oMath"))
        if not equations or (text and not EQUATION_NUMBER_RE.match(text)):
            return False
        for equation in equations:
            if latex := omml_to_latex(equation):
                self.outline.add(TextBlock(kind="equation", content=latex, format="latex"))
        return True

    def read_images(self, p: Paragraph) -> bool:
        """Save the paragraph's embedded images as figures. True when the paragraph holds any image."""
        rel_ids = [rid for blip in p._p.iter(qn("a:blip")) if (rid := blip.get(qn("r:embed")))]
        for rel_id in rel_ids:
            name = self.save_image(rel_id, len(self.figures) + 1)
            if name:
                figure = Figure(id=f"fig{len(self.figures) + 1}", file=name)
                self.figures.append(figure)
                self.outline.add(TextBlock(kind="figure_ref", content=figure.id))
            else:
                self.unsupported_images += 1
        return bool(rel_ids)

    def save_image(self, rel_id: str, index: int) -> str | None:
        part = self.document.part.related_parts.get(rel_id)
        ext = IMAGE_EXTENSIONS.get(getattr(part, "content_type", ""))
        if not ext:
            return None
        name = f"docx_figure_{index}{ext}"
        (self.figure_dir / name).write_bytes(part.blob)
        return name

    def add_caption(self, match: re.Match) -> bool:
        """Attach a "Figure N:" or "Table N:" caption. False when there is no figure waiting for one."""
        caption = latex_escape(match.group(3).strip())
        if match.group(1).lower().startswith("tab"):
            blocks = self.outline.blocks
            if self.tables and not self.tables[-1].caption and blocks and blocks[-1].kind == "table_ref":
                self.tables[-1].caption = caption  # caption below the table
            else:
                self.table_caption = caption  # caption above the table
            return True
        if self.figures and not self.figures[-1].caption:
            self.figures[-1].caption = caption
            return True
        return False

    def add_list_item(self, latex: str, style: str) -> None:
        """Add a "List Bullet" / "List Number" paragraph. "List Number 2" style items are sub-items.

        A bullet item right after numbered items (or the other way round) starts a new list.
        """
        ordered = "number" in style
        level = int(m.group(1)) if (m := re.search(r"(\d+)\s*$", style)) else 1
        blocks = self.outline.blocks
        if blocks and blocks[-1] is self.list_block and (level > 1 or self.list_block.ordered == ordered):
            self.list_entries.append((level, ordered, latex))
        else:
            self.list_block = TextBlock(kind="list", format="latex", ordered=ordered)
            self.list_entries = [(level, ordered, latex)]
            blocks.append(self.list_block)
        # Update the list in place: a plain list while all items share one level, nested LaTeX otherwise.
        block = self.list_block
        if len({lvl for lvl, _, _ in self.list_entries}) == 1:
            block.items = [text for _, _, text in self.list_entries]
        else:
            block.kind, block.items, block.content = "raw_latex", [], nested_list_latex(self.list_entries)


def _style_name(p: Paragraph) -> str:
    try:
        return (p.style.name or "").lower() if p.style is not None else ""
    except (KeyError, AttributeError):
        return ""


def _heading_level(p: Paragraph, text: str) -> int | None:
    """Level of a "Heading N" paragraph, or of a short numbered line set in bold or capitals."""
    style = _style_name(p)
    if m := re.match(r"heading\s*(\d)", style):
        return int(m.group(1))
    if style in {"title", "subtitle"} or len(text) >= 90:
        return None
    m = NUMBERED_HEADING.match(text)
    if not m:
        return None
    numbering = m.group(1)
    bold = bool(p.runs) and all(r.bold for r in p.runs if r.text.strip())
    if not (bold or text.isupper() or re.fullmatch(r"[IVX]+", numbering)):
        return None
    return numbering.count(".") + 1 if numbering[0].isdigit() else 1


def _marked_text(p: Paragraph) -> str:
    """The paragraph's text with superscript number runs as Unicode superscripts ("Chen Wei¹,²"), so author
    and affiliation markers stay recognisable."""
    parts = []
    for run in p.runs:
        text = run.text
        if run.font is not None and run.font.superscript and re.fullmatch(r"[\d,\s*]+", text):
            text = text.translate(TO_SUPERSCRIPT)
        parts.append(text)
    return normalize_space("".join(parts))


def _footnotes(document) -> dict[str, str]:
    """The document's footnotes as {id: LaTeX text}."""
    for rel in document.part.rels.values():
        if rel.reltype.endswith("/footnotes") and not rel.is_external:
            notes = parse_xml(rel.target_part.blob).iterchildren(qn("w:footnote"))
            return {note.get(qn("w:id")): latex_escape(normalize_space("".join(note.itertext(qn("w:t")))))
                    for note in notes if not note.get(qn("w:type"))}
    return {}


def _runs_to_latex(p: Paragraph, footnotes: dict[str, str]) -> str:
    """The paragraph as LaTeX: formatted runs (also those inside links and tracked insertions), footnotes and
    inline equations, in reading order."""
    parts = []
    for child in p._p.iterchildren():
        if child.tag == qn("m:oMath"):
            parts.append(f"${omml_to_latex(child)}$")
        elif child.tag == qn("w:r"):
            parts.append(_run_to_latex(Run(child, p), footnotes))
        elif child.tag in RUN_CONTAINERS:
            parts += [_run_to_latex(Run(r, p), footnotes) for r in child.iterchildren(qn("w:r"))]
    # Word often splits one bold phrase into several runs; join them back into one \textbf.
    latex = re.sub(r"\\textbf\{([^{}]*)\}\\textbf\{", r"\\textbf{\1", "".join(parts))
    return latex.strip()


def _run_to_latex(run: Run, footnotes: dict[str, str]) -> str:
    if (ref := run._r.find(qn("w:footnoteReference"))) is not None:
        note = footnotes.pop(ref.get(qn("w:id")), "")
        return f"\\footnote{{{note}}}" if note else ""
    if not run.text:
        return ""
    text = latex_escape(run.text)
    if run.bold and run.italic:
        text = f"\\textbf{{\\emph{{{text}}}}}"
    elif run.bold:
        text = f"\\textbf{{{text}}}"
    elif run.italic:
        text = f"\\emph{{{text}}}"
    if run.font is not None and run.font.superscript:
        text = f"\\textsuperscript{{{text}}}"
    elif run.font is not None and run.font.subscript:
        text = f"\\textsubscript{{{text}}}"
    return text
